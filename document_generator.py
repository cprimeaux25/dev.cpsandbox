"""Generate Word documents from extracted bid/RFP information."""

import io
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from bid_extractor import BidExtraction


def _set_cell_background(cell, color_hex: str):
    """Set background color for a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _add_heading(doc: Document, text: str, level: int = 1):
    """Add a styled heading."""
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.runs[0]
    if level == 1:
        run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
    elif level == 2:
        run.font.color.rgb = RGBColor(0x2E, 0x6D, 0xA4)
    return heading


def _add_info_table(doc: Document, rows: list[tuple[str, str]]):
    """Add a two-column information table."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, (label, value) in enumerate(rows):
        label_cell = table.rows[i].cells[0]
        value_cell = table.rows[i].cells[1]

        label_cell.width = Inches(2.2)
        value_cell.width = Inches(4.5)

        _set_cell_background(label_cell, "E8F0F8")

        label_para = label_cell.paragraphs[0]
        label_run = label_para.add_run(label)
        label_run.bold = True
        label_run.font.size = Pt(10)

        value_para = value_cell.paragraphs[0]
        value_run = value_para.add_run(value or "Not specified")
        value_run.font.size = Pt(10)
        if not value:
            value_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.add_paragraph()


def _add_bullet_list(doc: Document, items: list[str], style: str = "List Bullet"):
    """Add a bulleted list of items."""
    if not items:
        p = doc.add_paragraph("None specified", style="Normal")
        p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()
        return

    for item in items:
        if item.strip():
            p = doc.add_paragraph(style=style)
            run = p.add_run(item.strip())
            run.font.size = Pt(10)
    doc.add_paragraph()


def _add_section_divider(doc: Document):
    """Add a horizontal rule between sections."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E6DA4")
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_bid_document(bid: BidExtraction, output_filename: str = None) -> bytes:
    """Generate a Word document from extracted bid data and return bytes."""
    doc = Document()

    # Page margins
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    # Title block
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("BID / RFP SUBMISSION CHECKLIST")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)

    subtitle_para = doc.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_para.add_run(bid.rfp_title)
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor(0x2E, 0x6D, 0xA4)

    generated_para = doc.add_paragraph()
    generated_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated_run = generated_para.add_run(
        f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    )
    generated_run.font.size = Pt(9)
    generated_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph()
    _add_section_divider(doc)

    # ── Section 1: Project Overview ──────────────────────────────────────────
    _add_heading(doc, "1. Project Overview", level=1)

    overview_rows = [
        ("RFP / Bid Title", bid.rfp_title),
        ("Bid / RFP Number", bid.bid_number),
        ("Issuing Agency / Owner", bid.owner_agency),
        ("Project Location", bid.project_location),
        ("Engineer of Record", bid.engineer_name),
        ("Engineering Firm", bid.engineer_firm),
    ]
    _add_info_table(doc, overview_rows)
    _add_section_divider(doc)

    # ── Section 2: Critical Deadlines ────────────────────────────────────────
    _add_heading(doc, "2. Critical Deadlines", level=1)

    deadline_rows = [
        ("Bid Due Date", bid.due_date),
        ("Bid Due Time", bid.due_time),
        ("Deadline Notes", bid.due_date_notes),
        ("Questions Deadline", bid.questions_deadline),
        ("Pre-Bid Meeting", bid.pre_bid_meeting),
    ]
    _add_info_table(doc, deadline_rows)
    _add_section_divider(doc)

    # ── Section 3: Scope of Work ─────────────────────────────────────────────
    _add_heading(doc, "3. Scope of Work", level=1)
    if bid.scope_of_work:
        p = doc.add_paragraph(bid.scope_of_work)
        p.runs[0].font.size = Pt(10)
    else:
        p = doc.add_paragraph("Not specified")
        p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.runs[0].font.size = Pt(10)
    doc.add_paragraph()
    _add_section_divider(doc)

    # ── Section 4: Project Timeline ──────────────────────────────────────────
    _add_heading(doc, "4. Project Timeline & Schedule", level=1)
    if bid.project_timeline:
        p = doc.add_paragraph(bid.project_timeline)
        p.runs[0].font.size = Pt(10)
    else:
        p = doc.add_paragraph("Not specified")
        p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.runs[0].font.size = Pt(10)
    doc.add_paragraph()
    _add_section_divider(doc)

    # ── Section 5: Submission Requirements ───────────────────────────────────
    _add_heading(doc, "5. Submission Requirements", level=1)

    submission_rows = [
        ("Submission Method", bid.submission_method),
        ("Submission Address", bid.submission_address),
        ("Number of Copies", bid.number_of_copies),
    ]
    _add_info_table(doc, submission_rows)

    # Required Forms
    _add_heading(doc, "5.1 Required Forms", level=2)
    _add_bullet_list(doc, bid.forms_required)

    # Licenses & Certifications
    _add_heading(doc, "5.2 Required Licenses & Certifications", level=2)
    _add_bullet_list(doc, bid.licenses_required)

    # References
    _add_heading(doc, "5.3 References", level=2)
    if bid.references_required:
        p = doc.add_paragraph(bid.references_required)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()
    else:
        p = doc.add_paragraph("Not specified")
        p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()

    # Bonding
    _add_heading(doc, "5.4 Bonding Requirements", level=2)
    if bid.bonding_requirements:
        p = doc.add_paragraph(bid.bonding_requirements)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()
    else:
        p = doc.add_paragraph("Not specified")
        p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()

    # Notary
    _add_heading(doc, "5.5 Notarization Requirements", level=2)
    _add_bullet_list(doc, bid.notary_requirements)

    # Insurance
    _add_heading(doc, "5.6 Insurance Requirements", level=2)
    if bid.insurance_requirements:
        p = doc.add_paragraph(bid.insurance_requirements)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()
    else:
        p = doc.add_paragraph("Not specified")
        p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()

    _add_section_divider(doc)

    # ── Section 6: Qualifications ────────────────────────────────────────────
    _add_heading(doc, "6. Qualification Requirements", level=1)
    _add_bullet_list(doc, bid.qualifications)
    _add_section_divider(doc)

    # ── Section 7: Technical Information ─────────────────────────────────────
    _add_heading(doc, "7. Technical Information Required", level=1)
    _add_bullet_list(doc, bid.technical_information)
    _add_section_divider(doc)

    # ── Section 8: Approved Products ─────────────────────────────────────────
    _add_heading(doc, "8. Approved Products & Manufacturers", level=1)
    _add_bullet_list(doc, bid.approved_products)
    _add_section_divider(doc)

    # ── Section 9: Equipment & Product Specifications ─────────────────────────
    _add_heading(doc, "9. Equipment & Product Specifications", level=1)
    _add_bullet_list(doc, bid.equipment_specifications)
    _add_section_divider(doc)

    # ── Section 10: Water Meter Specifications ────────────────────────────────
    _add_heading(doc, "10. Water Meter Sizes & Quantities", level=1)
    if bid.water_meter_specs:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        header_cells = table.rows[0].cells
        headers = ["Meter Size", "Quantity", "Notes"]
        for i, header in enumerate(headers):
            _set_cell_background(header_cells[i], "1A3A5C")
            run = header_cells[i].paragraphs[0].add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(10)

        for spec in bid.water_meter_specs:
            row = table.add_row()
            row.cells[0].text = spec.size or ""
            row.cells[1].text = spec.quantity or "Not specified"
            row.cells[2].text = spec.notes or ""
            for cell in row.cells:
                cell.paragraphs[0].runs[0].font.size = Pt(10)

        doc.add_paragraph()
    else:
        p = doc.add_paragraph("No water meter specifications found in document")
        p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()

    _add_section_divider(doc)

    # ── Section 11: Additional Submission Items ───────────────────────────────
    if bid.additional_submission_items:
        _add_heading(doc, "11. Additional Submission Items", level=1)
        for item in bid.additional_submission_items:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(f"[{item.category}] {item.description}")
            run.font.size = Pt(10)
            if not item.required:
                run.italic = True
            if item.notes:
                note_run = p.add_run(f" — {item.notes}")
                note_run.font.size = Pt(9)
                note_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        doc.add_paragraph()
        _add_section_divider(doc)

    # ── Section 12: Special Conditions ───────────────────────────────────────
    _add_heading(doc, "12. Special Conditions & Requirements", level=1)

    special_rows = [
        ("Prevailing Wage", "Yes" if bid.prevailing_wage else ("No" if bid.prevailing_wage is False else "Not specified")),
        ("Minority Participation", bid.minority_requirements),
        ("Estimated Project Cost", bid.estimated_cost),
        ("Liquidated Damages", bid.liquidated_damages),
    ]
    _add_info_table(doc, special_rows)

    if bid.special_requirements:
        _add_heading(doc, "12.1 Additional Special Requirements", level=2)
        _add_bullet_list(doc, bid.special_requirements)

    _add_section_divider(doc)

    # Footer note
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run(
        "⚠ This document is AI-generated from the source RFP/bid document. "
        "Always verify requirements against the original document before submission."
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x80, 0x40, 0x00)
    footer_run.italic = True

    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
